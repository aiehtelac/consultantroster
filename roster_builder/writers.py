"""Write rostered teams into the two output workbooks.

Teams come from the solver's Result and, optionally, the week carried forward
from last year (carried.Carried): anything with .weeks and .ph.

Consultant Roster (the blank from template.py), per date row:
  C and F  the CA (weekday and weekend call)
  G-L      Passive 1-6, L being the R6 standby
  PH rows use the same columns with '-PH' after each short name.

Leave projection (a copy of the uploaded file), per person:
  CA       on all 7 days, Fri -> Thu
  R1-R6    on Sat and Sun, and on a PH day for that PH's team (a substitute
           takes the slot number of the person they replace)
Existing entries are never overwritten, except an 'OK' request or a code
judged not to be leave; clashes are returned so they can be checked.
"""

from __future__ import annotations

import datetime as dt

from openpyxl import load_workbook

from .leave import END_MARKER, REQUEST, column_dates, month_sheets, normalise
from .staff import Staff, clean
from .template import CA_COLS, DAY, FIRST_ROW, ROUND_COLS, as_date


def _merge(rosters):
    """All week teams and all PH teams across the given rosters."""
    weeks = {f: t for r in rosters for f, t in r.weeks.items()}
    ph = {d: t for r in rosters for d, t in r.ph.items()}
    return weeks, ph


def fill_roster(path, staff: list[Staff], *rosters) -> None:
    """Write the teams into the blank roster at `path`, in place."""
    short = {s.full_name: s.short_name for s in staff}
    weeks, ph = _merge(rosters)
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    for r in range(FIRST_ROW, ws.max_row + 1):
        day = as_date(ws.cell(r, 1).value)
        if day is None:  # end of the date block
            break
        is_ph = isinstance(ws.cell(r, 2).value, str)  # PH rows carry a label in B
        team = (ph if is_ph else weeks).get(day)
        if team is None:
            continue
        suffix = "-PH" if is_ph else ""
        for col, name in zip(CA_COLS + tuple(ROUND_COLS), [team.ca, team.ca, *team.rounds]):
            if name:  # a blank slot stays blank
                ws.cell(r, col).value = short.get(name, name) + suffix
    wb.save(path)


def duties(*rosters) -> dict[tuple[str, dt.date], str]:
    """(full name, date) -> code to enter in the leave projection."""
    weeks, ph = _merge(rosters)
    out = {}
    for fri, team in weeks.items():
        for i in range(7):
            out[team.ca, fri + i * DAY] = "CA"
        for slot, name in enumerate(team.rounds, start=1):
            out[name, fri + DAY] = out[name, fri + 2 * DAY] = f"R{slot}"
    for day, team in ph.items():
        for slot, name in enumerate(team.rounds, start=1):
            out[name, day] = f"R{slot}"
    return {key: code for key, code in out.items() if key[0]}  # drop blank slots


def fill_leave(src, dst, *rosters, not_leave=frozenset()) -> list[str]:
    """Copy the leave projection with duties entered. Returns any clashes.

    not_leave: (full name, date) entries judged not to be leave; a duty there
    replaces them instead of clashing.
    """
    todo = duties(*rosters)
    clashes = []
    wb = load_workbook(src)
    for ws in month_sheets(wb):
        dates = column_dates(ws)
        for row in ws.iter_rows(min_row=FIRST_ROW):
            name = clean(row[0].value)
            if name.upper().startswith(END_MARKER):
                break
            for cell in row[1:]:
                key = (name, dates.get(cell.column))
                code = todo.get(key)
                if code is None:
                    continue
                existing = normalise(cell.value)
                if existing and existing not in (REQUEST, code) and key not in not_leave:
                    clashes.append(f"{ws.title}!{cell.coordinate} {name}: has {cell.value!r}, wanted {code}")
                    continue
                cell.value = code
    wb.save(dst)
    return clashes
