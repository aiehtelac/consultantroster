"""Read the week carried forward from last year's Consultant Roster.

This year's first week row (the week containing 1 Jan) is last year's last
week, so its team - and the team on any PH inside it - is copied from last
year's roster rather than solved again. It is shown on the new roster but not
counted, and the solver keeps those people's next duty well spaced from it.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass

from openpyxl import load_workbook

from .model import Team
from .staff import Staff, clean
from .template import CA_COLS, DAY, FIRST_ROW, ROUND_COLS, as_date, fridays


@dataclass
class Carried:
    weeks: dict[dt.date, Team]  # the one week, keyed by its Friday
    ph: dict[dt.date, Team]  # any PH inside it
    notes: list[str]  # names matched loosely, or not at all

    @property
    def people(self) -> set[str]:
        teams = [*self.weeks.values(), *self.ph.values()]
        return {name for t in teams for name in [t.ca, *t.rounds] if name}


def read_carried(path, year: int, staff: list[Staff]) -> Carried:
    """Teams for the week starting on this year's first Friday, from last year's roster.

    Names are short names (with '-PH' on PH rows). Each becomes a full name;
    one that is no longer a short name is matched to the single full name it
    starts with (e.g. Chan Ah Meng ), and
    anything else is kept as written.
    """
    fri = fridays(year)[0]
    by_short = {s.short_name.lower(): s.full_name for s in staff}
    notes: dict[str, None] = {}  # ordered, without repeats

    def person(value) -> str | None:
        name = clean(value)
        if name.upper().endswith("-PH"):
            name = name[:-3].strip()
        if not name:
            return None
        if name.lower() in by_short:
            return by_short[name.lower()]
        hits = [s.full_name for s in staff if s.full_name.lower().startswith(name.lower())]
        if len(hits) == 1:
            notes[f"'{name}' taken as {hits[0]}"] = None
            return hits[0]
        notes[f"'{name}' is not in the staff list - copied as written, not counted"] = None
        return name

    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    weeks, ph = {}, {}
    for r in range(FIRST_ROW, ws.max_row + 1):
        day = as_date(ws.cell(r, 1).value)
        if day is None or not fri <= day <= fri + 6 * DAY:  # skip blank, text and other weeks' rows
            continue
        team = Team(person(ws.cell(r, CA_COLS[0]).value), [person(ws.cell(r, c).value) for c in ROUND_COLS])
        if isinstance(ws.cell(r, 2).value, str):  # PH rows carry a label in B
            ph[day] = team
        elif day == fri:
            weeks[day] = team
    if not weeks:
        raise ValueError(f"{path}: no week starting Fri {fri:%d %b %Y} - is it last year's roster?")
    return Carried(weeks, ph, list(notes))
