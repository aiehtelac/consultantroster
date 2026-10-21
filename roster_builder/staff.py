"""Read the staff list from StaffList_Targets.xlsx.

Targets sheet columns: A full name | B short name | C group | D active PAME calls
| E target calls. Call targets come from the group bands below, unless column E
holds a personal override such as '5-6' or '5' (store ranges as Text in Excel,
or it turns '5-6' into a date).
"""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass

from openpyxl import load_workbook

COL_NAME, COL_SHORT, COL_GROUP, COL_PAME, COL_TARGET = 1, 2, 3, 4, 5

# Weekend calls/rounds per year (min, max) for each group
BANDS = {
    "A": (6, 7),  # can do PAME calls
    "B": (6, 7),  # no PAME calls
    "C": (4, 5),  # no PAME calls
    "D": (1, 3),  # no PAME calls
    "E": (0, 0),  # no calls/rounds
}


def clean(value) -> str:
    """Normalise a cell to plain text: NBSP, odd unicode and stray spaces out."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    return " ".join(text.split())


@dataclass(frozen=True)
class Staff:
    full_name: str
    short_name: str
    group: str
    pame: bool
    tmin: int
    tmax: int

    @property
    def band(self) -> str:
        """Target as shown in the roster's Target column, e.g. '6-7'."""
        return f"{self.tmin}-{self.tmax}"


def parse_target(value, name: str) -> tuple[int, int] | None:
    """Column E override: '5-6' -> (5, 6), 5 -> (5, 5), blank -> None."""
    if isinstance(value, dt.datetime):
        raise ValueError(f"{name}: Excel turned the target into a date - format column E as Text and re-enter it")
    text = clean(value)
    if not text:
        return None
    found = re.fullmatch(r"(\d+)(?:\.0)?\s*(?:-\s*(\d+))?", text)
    if not found:
        raise ValueError(f"{name}: target {text!r} should look like '5-6' or '5'")
    low, high = int(found.group(1)), int(found.group(2) or found.group(1))
    if low > high:
        raise ValueError(f"{name}: target {text!r} has min above max")
    return low, high


def read_all(path) -> list[Staff]:
    """Everyone with a group, Group E included, in Targets-sheet order
    (service-grouped, matching the leave projection)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for row in wb["Targets"].iter_rows(min_row=2, values_only=True):
        name, group = clean(row[COL_NAME - 1]), clean(row[COL_GROUP - 1]).upper()
        if not name or not group:  # service header rows have no group
            continue
        if group not in BANDS:
            raise ValueError(f"{name}: unknown group {group!r} (expected one of {', '.join(BANDS)})")
        tmin, tmax = parse_target(row[COL_TARGET - 1], name) or BANDS[group]
        pame = clean(row[COL_PAME - 1]).upper() == "Y"
        out.append(Staff(name, clean(row[COL_SHORT - 1]), group, pame, tmin, tmax))
    wb.close()
    return out


def sheet_names(path) -> set[str]:
    """Everything in column A of the Targets sheet: people and service headings."""
    wb = load_workbook(path, data_only=True, read_only=True)
    names = {clean(row[0]) for row in wb["Targets"].iter_rows(min_row=2, max_col=1, values_only=True)}
    wb.close()
    return names - {""}


def unlisted(leave_names, targets_path) -> list[str]:
    """Names in the leave projection that are nowhere on the Targets sheet - most
    likely new staff, since the leave projection is usually the more up to date."""
    known = sheet_names(targets_path)
    return [n for n in leave_names if n not in known]


def rosterable(staff: list[Staff]) -> list[Staff]:
    """Those with calls to do (target above 0), after checking their names."""
    out = [s for s in staff if s.tmax > 0]
    missing = [s.full_name for s in out if not s.short_name]
    if missing:
        raise ValueError(f"no short name for: {', '.join(missing)}")
    for column, names in (("full", [s.full_name for s in out]), ("short", [s.short_name for s in out])):
        dupes = sorted(name for name, n in Counter(names).items() if n > 1)
        if dupes:
            raise ValueError(f"duplicate {column} names: {', '.join(dupes)}")
    return out


def read_staff(path) -> list[Staff]:
    """Rosterable staff from StaffList_Targets.xlsx (Group E left out)."""
    return rosterable(read_all(path))


def write_targets(src, dst, staff: list[Staff]) -> None:
    """Save edits into a copy of StaffList_Targets.xlsx, matching rows by full name.

    Column E gets a person's range only where it differs from their group's,
    stored as Text so Excel can't turn '5-6' into a date.
    """
    by_name = {s.full_name: s for s in staff}
    wb = load_workbook(src)
    for row in wb["Targets"].iter_rows(min_row=2):
        s = by_name.get(clean(row[COL_NAME - 1].value))
        if s is None:
            continue
        row[COL_SHORT - 1].value = s.short_name
        row[COL_GROUP - 1].value = s.group
        row[COL_PAME - 1].value = "Y" if s.pame else None
        row[COL_TARGET - 1].value = None if (s.tmin, s.tmax) == BANDS[s.group] else s.band
        row[COL_TARGET - 1].number_format = "@"
    wb.save(dst)
